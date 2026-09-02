"""Genera los .docx para subir a NotebookLM.

NotebookLM funciona mejor con documentos ACOTADOS: uno por tema, con títulos
reales y párrafos cortos. Un único archivo gigante hace que las respuestas
mezclen contextos que no tienen que ver entre sí. Por eso son diez.

    uv run python scripts_docx.py

Salida en docs/notebooklm/. Se puede volver a correr cuando cambien los
números: el contenido vive acá, no en los .docx.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

SALIDA = Path("docs/notebooklm")

# Cifras medidas en el proyecto. Se declaran una sola vez para que no se
# desincronicen entre documentos.
CIFRAS = {
    "productos": "149.595",
    "marcas": "46",
    "variantes": "206.082",
    "price_points": "247.366",
    "tests": "53",
    "adapters": "7",
    "avellaneda": "134.638",
    "marcas_prod": "14.957",
}


def doc_nuevo(titulo: str, bajada: str) -> Document:
    d = Document()
    est = d.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(11)

    h = d.add_heading(titulo, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = d.add_paragraph()
    r = p.add_run(bajada)
    r.italic = True
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return d


def h1(d: Document, t: str) -> None:
    d.add_heading(t, level=1)


def h2(d: Document, t: str) -> None:
    d.add_heading(t, level=2)


def p(d: Document, t: str) -> None:
    d.add_paragraph(t)


def bullets(d: Document, items: list[str]) -> None:
    for i in items:
        d.add_paragraph(i, style="List Bullet")


def numerada(d: Document, items: list[str]) -> None:
    for i in items:
        d.add_paragraph(i, style="List Number")


def tabla(d: Document, cabeceras: list[str], filas: list[list[str]]) -> None:
    t = d.add_table(rows=1, cols=len(cabeceras))
    t.style = "Light Grid Accent 1"
    for celda, texto in zip(t.rows[0].cells, cabeceras):
        celda.text = texto
        for par in celda.paragraphs:
            for run in par.runs:
                run.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for celda, texto in zip(celdas, fila):
            celda.text = texto
    d.add_paragraph()


def cita(d: Document, t: str) -> None:
    par = d.add_paragraph()
    run = par.add_run(t)
    run.italic = True
    par.paragraph_format.left_indent = Pt(24)


# ─────────────────────────────────────────────────────────────────────────
# 01 · Qué es FadIA
# ─────────────────────────────────────────────────────────────────────────
def doc01() -> Document:
    d = doc_nuevo(
        "FadIA · Qué es y qué problema resuelve",
        "Buscador conversacional de moda argentina sobre catálogos scrapeados.",
    )

    h1(d, "El problema")
    p(d, "Comprar ropa en Argentina implica recorrer decenas de sitios que no "
         "hablan entre sí. Cada marca tiene su ecommerce, cada mayorista su "
         "vidriera, y ninguno permite preguntar en lenguaje natural. Buscar "
         "«un vestido negro para un casamiento» exige traducir esa intención a "
         "filtros de categoría, color y talle, sitio por sitio.")
    p(d, "Del lado mayorista el problema es distinto y peor: Av. Avellaneda "
         "concentra miles de locales cuyo catálogo online existe pero está "
         "disperso, sin precio por mayor visible ni compra mínima declarada de "
         "forma consistente.")

    h1(d, "Qué hace FadIA")
    p(d, "Scrapea las vidrieras públicas de marcas y mayoristas, normaliza todo "
         "a un esquema canónico único, y expone una interfaz conversacional en "
         "castellano. El usuario pregunta y refina en lenguaje natural; el "
         "sistema devuelve productos reales con precio, talles y link a la "
         "tienda.")

    h2(d, "Dos modos de compra")
    tabla(d, ["Modo", "Productos", "Qué le importa al usuario"], [
        ["Marcas", CIFRAS["marcas_prod"], "Talle, marca, ocasión, caída de la prenda"],
        ["Por mayor", CIFRAS["avellaneda"], "Precio por mayor, curva mínima, dirección del local"],
    ])

    h1(d, "La decisión de arquitectura central")
    p(d, "Postgres resuelve la búsqueda. El modelo de lenguaje solo redacta.")
    p(d, "El diseño intuitivo sería un «wrapper»: el modelo consulta la base, "
         "elige productos y contesta. Ese patrón se descartó. En FadIA el "
         "modelo recibe las filas que Postgres ya eligió, como contexto de solo "
         "lectura. No puede elegir productos ni inventar precios.")

    h2(d, "Por qué importa, con un caso real")
    p(d, "En una prueba la recuperación falló y al modelo le llegaron "
         "cosméticos cuando el usuario pedía remeras. La respuesta fue:")
    cita(d, "«No hay productos de indumentaria en el contexto que me pasaste.»")
    p(d, "Si el modelo hubiera podido elegir, habría rellenado el hueco con "
         "productos plausibles e inexistentes. Un buscador de compras que "
         "inventa un precio o un talle es peor que uno que no encuentra nada.")

    h1(d, "Escala actual")
    tabla(d, ["Métrica", "Valor"], [
        ["Productos", CIFRAS["productos"]],
        ["Marcas distintas", CIFRAS["marcas"]],
        ["Variantes (talle × color)", CIFRAS["variantes"]],
        ["Puntos de precio históricos", CIFRAS["price_points"]],
        ["Adapters de plataforma", CIFRAS["adapters"]],
        ["Tests de regresión", CIFRAS["tests"]],
        ["Productos con embedding", "100 %"],
    ])

    h1(d, "Restricción de diseño: todo corre local")
    p(d, "No hay API paga ni servicio externo. El modelo de chat y el de "
         "embeddings corren en la máquina del usuario mediante LM Studio. "
         "Consecuencias: costo por consulta cero, el catálogo no sale de la "
         "máquina, y funciona sin conexión una vez cargados los datos.")
    p(d, "Está probado en una MacBook Air M2 de 16 GB.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 02 · Arquitectura
# ─────────────────────────────────────────────────────────────────────────
def doc02() -> Document:
    d = doc_nuevo(
        "FadIA · Arquitectura del sistema",
        "Dos mitades que se cruzan en la base: ingesta y consulta.",
    )

    h1(d, "Mitad de ingesta")
    numerada(d, [
        "Adapters. Siete, uno por plataforma de ecommerce. Cada uno sabe dónde "
        "vive la verdad en esa plataforma y devuelve objetos Product canónicos.",
        "Crawler adaptativo. Control de caudal AIMD (aumento aditivo, "
        "reducción multiplicativa) con guardia de latencia. Respeta robots.txt "
        "y arranca en 1 req/s.",
        "Arnés de validación. Compara la forma del resultado contra lo que ya "
        "se sabe del sitio y bloquea el lote antes de escribir en la base.",
        "Persistencia. Postgres 17 con pgvector y pg_trgm.",
    ])

    h2(d, "Detección de plataforma en cadena")
    p(d, "Ante un dominio nuevo, el sistema prueba en orden: host conocido, "
         "firma en el HTML, y sondeo de endpoints. El sondeo por endpoint "
         "resultó ser el más confiable: Zara, por ejemplo, no menciona la "
         "palabra «zara» ni una vez en su home.")

    h1(d, "Mitad de consulta")
    numerada(d, [
        "La pregunta y el hilo previo se convierten en filtros duros por "
        "reglas: categoría, color, talle, tela, detalle, precio.",
        "El mismo texto acumulado se convierte en un vector de 768 dimensiones.",
        "Postgres combina filtros duros y distancia coseno, y devuelve las "
        "filas en unos 40 ms.",
        "Los resultados van a la interfaz de inmediato.",
        "En paralelo, el modelo recibe esas filas y redacta.",
    ])

    h2(d, "El orden importa para la percepción")
    p(d, "La grilla de productos llega a los 40 ms. El texto del modelo empieza "
         "a los 200 ms y tarda alrededor de 1,4 segundos en completarse. El "
         "usuario nunca mira una pantalla vacía: ve los productos primero y el "
         "comentario después.")

    h1(d, "Componentes y responsabilidades")
    tabla(d, ["Componente", "Responsabilidad"], [
        ["src/fadia/adapters/", "Un archivo por plataforma; todos devuelven Product"],
        ["src/fadia/normalize/", "money, size, color, taxonomy, text"],
        ["src/fadia/validate.py", "El arnés que bloquea lotes malos"],
        ["src/fadiaapi/search.py", "Filtros, facetas y consultas SQL"],
        ["src/fadiaapi/main.py", "Endpoints HTTP y prompt del sistema"],
        ["src/fadiaapi/llm.py", "Único archivo que habla con el runtime del modelo"],
        ["web/index.html", "La interfaz completa, sin build ni dependencias"],
    ])

    h2(d, "El front es un solo archivo")
    p(d, "web/index.html contiene el HTML, el CSS y el JavaScript. Es "
         "deliberado: no hay bundler, no hay node_modules, no hay paso de "
         "compilación. Se edita y se recarga.")

    h1(d, "Despliegue")
    p(d, "Tres contenedores Docker: base de datos, API y servidor web. Los "
         "modelos de lenguaje corren en el host, no en Docker, porque un "
         "contenedor Linux en Mac no tiene acceso a la GPU. Los servicios les "
         "hablan por host.docker.internal.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 03 · Adapters
# ─────────────────────────────────────────────────────────────────────────
def doc03() -> Document:
    d = doc_nuevo(
        "FadIA · Los siete adapters",
        "Uno por plataforma de ecommerce, no por tienda.",
    )

    h1(d, "El principio")
    p(d, "El tema visual cambia en cada tienda; el endpoint no. Escribir un "
         "parser por tienda es trabajo lineal que se rompe en cada rediseño. "
         "Escribir uno por plataforma es trabajo constante que sobrevive a los "
         "rediseños.")
    p(d, "La hipótesis se midió: de diez tiendas Tiendanube evaluadas, nueve "
         "funcionaron con el mismo adapter sin cambios.")

    h1(d, "Tabla comparativa")
    tabla(d, ["Adapter", "Fuente de verdad", "Talles", "Stock"], [
        ["tiendanube", "atributo data-variants", "sí", "sí"],
        ["avellaneda", "sitemap + JSON-LD", "no", "no"],
        ["shopify", "/products.json", "sí", "no"],
        ["vtex", "catalog_system", "sí", "sí"],
        ["magento", "/graphql", "sí", "no"],
        ["zara", "categories?ajax=true", "no", "no"],
        ["generico", "JSON-LD / microdata", "no", "no"],
    ])

    h2(d, "Los «no» no son fallas")
    p(d, "Shopify y Magento no publican cantidad de stock: uno manda null, el "
         "otro un enum. Zara no expone talles en el listado. En esos casos el "
         "campo queda en None y no en cero. Un dato ausente es información; un "
         "dato falso es una mentira que después nadie detecta.")

    h1(d, "La trampa de cada plataforma")

    h2(d, "Tiendanube")
    p(d, "El JSON-LD publica el precio de LISTA, no el de venta. Un vestido "
         "aparecía a $36.500 cuando se vendía a $23.725: 54 % de inflación. La "
         "verdad está en el atributo data-variants del formulario de carrito.")
    p(d, "Además, algunas tiendas sirven páginas de «producto no encontrado» "
         "con HTTP 200. Hay una función dedicada a detectarlas.")

    h2(d, "Avellaneda")
    p(d, "Los listados por rubro se randomizan en cada request, así que "
         "paginar no garantiza cobertura. La enumeración completa se resuelve "
         "por el índice de sitemaps: 14 archivos que cubren el 100 % del "
         "marketplace.")

    h2(d, "Shopify")
    p(d, "El campo compare_at_price llega como la cadena «0.00» y no como "
         "null cuando no hay descuento. Tratarlo como número da un descuento "
         "del 100 %.")

    h2(d, "VTEX")
    p(d, "El total de resultados no viene en el cuerpo sino en un header. Y "
         "en varias tiendas la API pública no responde en el dominio "
         "comercial: hay que caer al dominio interno vtexcommercestable.com.br.")

    h2(d, "Magento")
    p(d, "El campo uid no existe antes de la versión 2.4.2 y su presencia "
         "tumba la query GraphQL entera. Hay una query alternativa de "
         "reserva sin ese campo.")
    p(d, "Aparte, las imágenes se sirven por una ruta de caché que devuelve el "
         "logo de Magento —un cuadrado de 262×262— cuando la variante de "
         "redimensionado nunca se generó. No da 404: da HTTP 200 con la imagen "
         "equivocada.")

    h2(d, "Zara")
    p(d, "La home no contiene la palabra «zara», así que la detección por "
         "firma de HTML falla. Se detecta por host. Los precios ya vienen en "
         "centavos, y las URLs de imagen hay que tomarlas del payload: "
         "construirlas a mano da 404.")

    h2(d, "Genérico")
    p(d, "La premisa era que JSON-LD cubriría cualquier sitio con SEO "
         "razonable. Se midió y es falsa para la cola larga argentina: de 25 "
         "marcas sin adapter dedicado, solo 2 publican el tipo Product, y una "
         "sola incluye precio.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 04 · Datos y normalización
# ─────────────────────────────────────────────────────────────────────────
def doc04() -> Document:
    d = doc_nuevo(
        "FadIA · Modelo de datos y normalización",
        "El esquema canónico y las decisiones de normalización.",
    )

    h1(d, "Jerarquía")
    p(d, "Producto → variantes → puntos de precio. Un producto agrupa lo que "
         "el usuario percibe como una prenda. Una variante es una combinación "
         "concreta de talle y color, con su propio precio y disponibilidad. Un "
         "punto de precio es una observación fechada.")
    p(d, f"Hoy: {CIFRAS['productos']} productos, {CIFRAS['variantes']} "
         f"variantes y {CIFRAS['price_points']} puntos de precio.")

    h2(d, "La serie temporal es el activo")
    p(d, "El histórico de precios no estaba en el diseño original y resultó ser "
         "lo más difícil de replicar para un competidor: se construye "
         "observando todos los días, no se puede comprar ni scrapear de una vez.")

    h1(d, "Sistemas de talle")
    p(d, "Un solo campo de texto no alcanza. Se clasifican en sistemas:")
    tabla(d, ["Sistema", "Ejemplo", "Origen del caso"], [
        ["ALPHA", "S, M, L, XL", "el más común"],
        ["AR_NUMERIC", "38, 40, 42", "talles argentinos de mujer"],
        ["SHOE_AR", "36, 37, 38", "calzado"],
        ["KIDS_NUMERIC", "4, 6, 8, 10, 12", "ropa de nena"],
        ["DIMENSION", "40x40", "productos con medidas"],
        ["ONE_SIZE", "único", "talle único"],
    ])

    h2(d, "Por qué existe KIDS_NUMERIC")
    p(d, "Una tienda guardaba los talles de nena (4, 6, 8, 10, 12) en el campo "
         "de color. El arnés lo detectó por una regla de plausibilidad: un "
         "conjunto de colores donde todos los valores son numéricos y "
         "consecutivos no es un conjunto de colores.")

    h1(d, "Taxonomía de categorías")
    p(d, "Diccionario de patrones evaluados en orden, contra el breadcrumb "
         "primero y el título después. El breadcrumb pesa más porque expresa "
         "la intención de la tienda.")

    h2(d, "El caso calza / calzado")
    p(d, "El patrón «calza» —la prenda— hacía match por prefijo con «calzado», "
         "y todo el calzado terminaba clasificado como pantalón. La corrección "
         "fue exigir palabra completa con plural opcional. Es un ejemplo de "
         "cómo una decisión de una línea contamina decenas de miles de filas.")

    h2(d, "Español peninsular")
    p(d, "Varias marcas internacionales publican en español de España. El "
         "diccionario incluye chaqueta, jersey, gorro, sudadera, camiseta y "
         "cazadora junto a los términos rioplatenses.")

    h1(d, "Inferencia de género")
    p(d, "Se resuelve en cadena: primero por categoría —vestidos, polleras y "
         "lencería vienen del lado femenino en los catálogos relevados—, "
         "después por dominancia de marca. Esto bajó el «sin género» del 42 % "
         "al 13,5 %.")

    h1(d, "Vocabulario textil argentino")
    p(d, "Las facetas de refinamiento no salen de una lista fija sino de los "
         "títulos reales del catálogo:")
    tabla(d, ["Término", "Prendas"], [
        ["morley", "4.903"],
        ["frisa", "3.365"],
        ["bengalina", "1.216"],
    ])
    p(d, "Ninguno tiene equivalente directo en el inglés de moda. Es "
         "vocabulario que ningún buscador extranjero puede tener sin haber "
         "leído catálogos argentinos.")

    h1(d, "Deuda conocida: la normalización se hornea al scrapear")
    p(d, "Las funciones de clasificación se llaman solo en los adapters. El "
         "archivo intermedio guarda el resultado ya cocinado y la ingesta lo "
         "copia sin recalcular. Consecuencia: arreglar el clasificador no "
         "corrige lo ya guardado.")
    bullets(d, [
        "3.514 productos de calzado siguen clasificados como pantalón, de "
        "5.361 en total. El clasificador actual los resuelve bien; las filas "
        "son anteriores a la corrección.",
        "2.772 variantes en 219 productos tienen talles con relleno de ceros "
        "(S00, M00, XL0) porque el normalizador no reconoce ese formato.",
    ])
    p(d, "Ambas se arreglan con un pase de reclasificación offline: el archivo "
         "intermedio conserva título y breadcrumb, que es todo lo que el "
         "clasificador necesita. No hace falta volver a scrapear.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 05 · Búsqueda
# ─────────────────────────────────────────────────────────────────────────
def doc05() -> Document:
    d = doc_nuevo(
        "FadIA · Búsqueda semántica y memoria del hilo",
        "Cómo se combinan filtros exactos, vectores y contexto conversacional.",
    )

    h1(d, "Dos mecanismos distintos")
    p(d, "No hay que confundirlos, porque fallan distinto y se arreglan "
         "distinto.")
    tabla(d, ["Mecanismo", "Qué captura", "Cómo persiste en el hilo"], [
        ["Filtros duros", "categoría, color, talle, tela, detalle, precio",
         "se acumulan turno a turno; lo nuevo pisa lo viejo"],
        ["Vector semántico", "todo lo demás: ocasión, registro, uso",
         "se arma con los últimos 3 turnos más el actual"],
    ])

    h1(d, "El caso que obligó a cambiar el diseño")
    p(d, "Guion de tres turnos: «un vestido para un casamiento», después «en "
         "verde», después «algo más largo».")
    p(d, "Con el vector armado solo con el último mensaje, el tercer turno "
         "devolvía abrigos. Los filtros duros mantenían la búsqueda dentro de "
         "vestidos verdes, pero el ORDEN dentro de ese conjunto lo decidía un "
         "vector que significa «tapado largo».")
    p(d, "La razón: «casamiento» es una OCASIÓN. No es categoría, ni color, ni "
         "talle, ni tela. No es filtrable, y por lo tanto vive únicamente en el "
         "embedding. Al rearmar el vector desde cero en cada turno, la ocasión "
         "se evaporaba.")

    h2(d, "Comprobación A/B")
    tabla(d, ["Texto que genera el vector", "Resultados"], [
        ["«algo mas largo»", "ABRIGO APATITO, ABRIGO CORTO, ABRIGO CAPA"],
        ["«un vestido para un casamiento en verde, algo mas largo»",
         "VESTIDO LAPACHO, Vestido Sueil, Vestido Mocri"],
    ])

    h2(d, "La regla de peso")
    p(d, "El mensaje actual se repite hasta pesar más que cualquier turno "
         "anterior individual, no más que la suma de todos. La regla estricta "
         "obligaba a descartar los turnos viejos para que el actual ganara, "
         "que es justo lo contrario de lo que se busca. Con la regla adoptada "
         "el mensaje actual queda en torno al 53 % del texto: domina, y el hilo "
         "sobrevive.")
    p(d, "Un agravante que explicaba la sensación de incoherencia: el modelo SÍ "
         "recibe el historial de la conversación. Se acordaba del casamiento y "
         "hablaba del casamiento, pero se le pasaban abrigos. Recordaba la "
         "charla y no los productos.")

    h1(d, "pgvector y el índice HNSW")
    p(d, "Embeddings de 768 dimensiones indexados con HNSW, más pg_trgm como "
         "respaldo por texto cuando el modelo no responde.")

    h2(d, "hnsw.iterative_scan no es opcional")
    p(d, "HNSW es un índice aproximado que POST-filtra: busca los vecinos más "
         "cercanos y recién después aplica el WHERE. Con un filtro selectivo "
         "devolvía 0 filas teniendo 1.938 coincidencias reales, sin error ni "
         "advertencia.")
    p(d, "Peor: la relajación automática de filtros tapaba el síntoma "
         "devolviendo maquillaje cuando se pedía ropa. La corrección es "
         "establecer hnsw.iterative_scan en relaxed_order.")

    h1(d, "Relajación de filtros")
    p(d, "Cuando no hay coincidencias exactas, los filtros se aflojan en un "
         "orden deliberado: primero detalle, después tela, color, talle, "
         "categoría, tienda y precio mínimo. Lo que se afloja se informa al "
         "usuario con un chip, y el modelo lo menciona en una línea antes de "
         "recomendar.")

    h1(d, "Facetas por eje")
    p(d, "Las sugerencias de refinamiento se agrupan en tres ejes —tela, "
         "silueta, detalle— y solo se muestran los términos que efectivamente "
         "discriminan dentro de los resultados actuales. Los ejes ya "
         "preguntados en el hilo se saltean, para que la conversación avance en "
         "vez de repetirse.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 06 · El modelo
# ─────────────────────────────────────────────────────────────────────────
def doc06() -> Document:
    d = doc_nuevo(
        "FadIA · El modelo de lenguaje",
        "Qué hace, qué no puede hacer, y cómo se lo hizo rápido.",
    )

    h1(d, "Posición en el sistema")
    p(d, "El modelo está AL COSTADO del camino de la búsqueda, no dentro. "
         "Recibe las filas que Postgres eligió, en un bloque de contexto de "
         "solo lectura, y escribe sobre ellas. No consulta la base, no elige "
         "productos, no calcula precios.")
    p(d, "Se usa Gemma 4 (variante e4b) corriendo local en LM Studio. La "
         "elección no fue por calidad máxima sino por el conjunto: corre en una "
         "laptop, no cuesta por consulta, y el catálogo no sale de la máquina.")

    h1(d, "Optimización de latencia: de 30 s a 1,4 s")

    h2(d, "Hallazgo 1 · El razonamiento no aportaba")
    p(d, "Gemma 4 es un modelo de razonamiento. En una prueba gastó 728 de 885 "
         "tokens pensando antes de escribir una línea. En este caso de uso el "
         "trabajo difícil —elegir qué productos— ya lo hizo Postgres, así que "
         "esa cadena no agrega nada. Apagarla con reasoning_effort en «none» "
         "llevó la respuesta de 29 s a 6 s.")

    h2(d, "Hallazgo 2 · Los modelos se desalojaban entre sí")
    p(d, "El culpable grande no era obvio. El modelo de embeddings desalojaba "
         "al de chat de la memoria en cada consulta. Una llamada de embeddings "
         "de 400 ms terminaba costando 3,8 s de recarga, y el primer token "
         "pasaba de 0,08 s a 4,4 s.")
    p(d, "La solución es cargar ambos con un TTL largo para que queden "
         "residentes.")

    h2(d, "Hallazgo 3 · El presupuesto de tokens tiene que ser holgado")
    p(d, "Con max_tokens en 120 el modelo devolvía una cadena vacía y 117 "
         "tokens de razonamiento: se gastaba el presupuesto pensando. El valor "
         "por defecto es 1600.")

    h1(d, "Streaming")
    p(d, "La respuesta se emite por SSE. El razonamiento viaja en un canal "
         "separado del contenido, lo que permite mostrarlo plegado en vez de "
         "mezclarlo con la respuesta.")

    h1(d, "El tono")
    p(d, "El registro es editorial, como una nota de pasarela: silueta, caída, "
         "textura, paleta, ocasión. Frases cortas y afirmativas.")
    p(d, "El prompt prohíbe explícitamente el registro coloquial rioplatense "
         "—«che», «te tiro», «copado», «onda»— porque en un contexto de moda "
         "leía como forzado. Es una corrección que salió del uso, no del "
         "diseño inicial.")

    h1(d, "Capacidad multimodal, verificada")
    p(d, "Gemma 4 acepta imágenes. Probado con una foto real del catálogo, "
         "respondió: «Es una sudadera con capucha (hoodie) de color negro, "
         "aparentemente confeccionada en una tela de felpa o algodón grueso. "
         "Presenta una silueta holgada y casual.» Correcto, en 4,7 segundos "
         "incluyendo la descarga.")
    p(d, "Detalle revelador: dijo «sudadera», no «buzo». El vocabulario visual "
         "del modelo es peninsular, igual que pasaba con la taxonomía.")

    h1(d, "Oportunidad: embeddings de imagen")
    p(d, "El modelo de embeddings actual es solo texto. Su par de visión está "
         "entrenado para compartir el mismo espacio latente de 768 dimensiones, "
         "lo que permitiría embeber las fotos y buscarlas con los embeddings de "
         "texto existentes: misma columna, mismo índice, sin recalcular el "
         "corpus.")
    p(d, "Dónde pagaría más: 43.383 productos —el 29 %— no tienen descripción "
         "usable y su vector es prácticamente solo el título. Para esos, la "
         "foto es la única señal real que existe.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 07 · Incidentes
# ─────────────────────────────────────────────────────────────────────────
def doc07() -> Document:
    d = doc_nuevo(
        "FadIA · Catálogo de fallas silenciosas",
        "Ninguno de los peores incidentes tiró una excepción.",
    )

    h1(d, "El patrón")
    p(d, "Todos los incidentes graves de este proyecto tienen la misma forma: "
         "HTTP 200, sin excepción, sin log de error, con datos incorrectos. Un "
         "sistema que falla ruidosamente es fácil; uno que falla en silencio "
         "corrompe la base durante semanas.")
    p(d, "Por eso la pieza más usada del proyecto no es un parser sino el arnés "
         "que compara la FORMA del resultado contra lo que ya se sabe del sitio.")

    h1(d, "Incidentes de ingesta")

    h2(d, "Precios inflados un 54 %")
    p(d, "El JSON-LD de Tiendanube publica el precio de lista. Un vestido "
         "figuraba a $36.500 cuando se vendía a $23.725. Sin comparar contra "
         "otra fuente, el dato es indistinguible de uno correcto.")

    h2(d, "Talles de nena guardados como colores")
    p(d, "Los valores 4, 6, 8, 10 y 12 terminaron en el campo de color. "
         "Detectado por una regla de plausibilidad sobre el conjunto.")

    h2(d, "600 páginas descargadas sin decodificar")
    p(d, "Se enviaba Accept-Encoding: br sin tener brotli instalado. "
         "Resultado: 600 respuestas HTTP 200, cero productos extraídos, ningún "
         "error en ningún lado.")

    h2(d, "Páginas 404 servidas con HTTP 200")
    p(d, "Una tienda devolvía su página de «producto no encontrado» con código "
         "200. El arnés lo detectó como duplicados: cientos de productos con "
         "contenido idéntico.")

    h2(d, "El placeholder de Magento")
    p(d, "Las imágenes se piden por una ruta de caché. Si esa variante de "
         "redimensionado nunca se generó, Magento devuelve su logo —262×262— "
         "con HTTP 200. Una tienda entera tenía el catálogo así. La foto "
         "original existe en la ruta sin el segmento de caché.")

    h1(d, "Incidentes de búsqueda")

    h2(d, "Cero filas de 1.938 coincidencias")
    p(d, "El post-filtrado de HNSW descrito en el documento de búsqueda. "
         "Silencioso, y tapado por la relajación automática de filtros.")

    h2(d, "La ocasión que se evaporaba")
    p(d, "El vector se rearmaba con el último mensaje y perdía todo lo no "
         "filtrable del hilo. Al tercer turno, «algo más largo» daba abrigos.")

    h1(d, "Incidentes de interfaz")

    h2(d, "El modal invisible que se comía todos los toques")
    p(d, "La regla [hidden]{display:none} es del navegador y pierde contra "
         "cualquier display de autor. El modal cerrado quedaba en pantalla con "
         "opacidad cero, capturando cada toque. En escritorio no se notaba "
         "porque, vacío, colapsaba a altura cero; en mobile un media query le "
         "fijaba altura completa y bloqueaba la pantalla entera.")

    h2(d, "Los controles que se plegaban antes del clic")
    p(d, "El despliegue dependía de :hover sobre un contenedor con "
         "pointer-events desactivado. Entre el input y los chips había una "
         "banda muerta: al cruzarla el contenedor perdía el hover y los "
         "controles desaparecían a mitad de camino. Se reemplazó por cercanía "
         "geométrica del puntero.")

    h2(d, "El anillo que no se veía")
    p(d, "La tarjeta tiene overflow oculto, que recorta cualquier sombra hacia "
         "afuera. El resaltado se veía como un borde tenue hasta que se pasó a "
         "sombra interior.")

    h1(d, "Falsos positivos al medir")
    p(d, "Tres veces pareció haber un bug de cascada CSS que no existía: las "
         "transiciones se congelan cuando la pestaña del navegador no tiene "
         "foco, y getComputedStyle devuelve el valor anterior. Se descarta "
         "quitando la transición y volviendo a leer.")
    p(d, "Lección transversal: la mitad de los bugs de interfaz eran invisibles "
         "leyendo el código y obvios midiendo en el navegador.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 08 · Interfaz
# ─────────────────────────────────────────────────────────────────────────
def doc08() -> Document:
    d = doc_nuevo(
        "FadIA · Diseño de la interfaz",
        "El chat como centro, y por qué cada animación existe.",
    )

    h1(d, "Principio: mobile-first y sin build")
    p(d, "Toda la interfaz vive en un archivo HTML con CSS y JS embebidos. Se "
         "diseñó primero para 320 px de ancho y después se expandió.")

    h1(d, "El chat es el centro")
    p(d, "En la pantalla inicial el campo de texto va inmediatamente debajo de "
         "las sugerencias, no clavado al pie: es lo primero que hay que mirar. "
         "Cuando empieza la conversación pasa a ser una barra fija, que es lo "
         "útil cuando hay resultados arriba.")

    h2(d, "Del tooltip al resultado")
    p(d, "Tocar una sugerencia no dispara la búsqueda directamente. La "
         "secuencia es: salto al chat, autoescritura del texto ahí, envío. "
         "Así el texto queda editable antes de mandarse y se entiende que todo "
         "pasa por el chat.")

    h1(d, "Las animaciones y su justificación")

    h2(d, "El contorno de lectura")
    p(d, "Mientras el modelo redacta, cada tarjeta se enciende con un contorno "
         "animado, escalonado 110 ms entre una y otra. No es decorativo: "
         "mientras el contorno gira, el modelo está efectivamente leyendo esas "
         "filas. Se eligió ese momento porque es el único tramo con espera "
         "real: la grilla llega a los 40 ms y el texto tarda 1,4 s.")

    h2(d, "La salida de impresora")
    p(d, "Los resultados no aparecen: salen. Un recorte los descubre de arriba "
         "hacia abajo con una línea fina haciendo de borde del papel.")

    h2(d, "El latido de la prenda citada")
    p(d, "Cuando la respuesta dice «la [3] destaca por el detalle del tejido», "
         "esa tarjeta se enciende con un anillo durante dos segundos. El pulso "
         "se dispara por cita y no al final: como el texto llega en streaming, "
         "cada prenda se destaca en el momento en que el modelo la nombra.")

    h1(d, "El modal de producto")
    p(d, "Modal centrado de dos columnas: fotos con miniaturas a la izquierda, "
         "ficha a la derecha. En mobile es una sola columna a pantalla completa.")
    bullets(d, [
        "Selección de color y talle donde la disponibilidad de talle depende "
        "del color elegido, no del producto en general.",
        "Carrusel con avance automático cada 2 segundos, flechas sutiles, y "
        "pausa al hacer clic sobre la foto.",
        "Link directo a la variante exacta, con el identificador de la opción "
        "en la URL.",
        "Solapa de envíos que muestra procedencia real y no inventa políticas.",
    ])

    h2(d, "Un caso extremo que obligó a decidir")
    p(d, "Un producto tenía 21 colores y casi todos agotados. Sin tope, el "
         "bloque de color empujaba los talles fuera de la pantalla. Se ordenan "
         "los disponibles primero y se limita la altura con scroll.")

    h1(d, "Rendimiento de imágenes")
    p(d, "Las tiendas publican la foto de catálogo entera —hasta 3000×3000— "
         "para pintarla en una tarjeta de 160 px. Dos CDN aceptan que se les "
         "pida el tamaño, y son justo los lentos.")
    tabla(d, ["CDN", "Original", "Redimensionada"], [
        ["VTEX", "2891 ms", "1040 ms"],
        ["Shopify", "804 ms", "431 ms"],
        ["Avellaneda", "249 ms", "ignora el parámetro"],
    ])
    p(d, "En estado estable, con el CDN caliente, la diferencia baja al 26 %. "
         "Lo que resuelve el parpadeo perceptible es otra cosa: la galería del "
         "modal no usa carga diferida y precarga la siguiente foto.")

    h1(d, "Accesibilidad y movimiento")
    p(d, "Todas las animaciones respetan prefers-reduced-motion. Las que "
         "dependen de hover tienen equivalente táctil: los controles plegados "
         "se despliegan manteniendo apretado el campo de texto.")
    return d


# ─────────────────────────────────────────────────────────────────────────
# 09 · Operación
# ─────────────────────────────────────────────────────────────────────────
def doc09() -> Document:
    d = doc_nuevo(
        "FadIA · Guía de instalación y operación",
        "Cómo levantarlo, cargarlo y reiniciarlo.",
    )

    h1(d, "Requisitos")
    tabla(d, ["Herramienta", "Necesidad", "Para qué"], [
        ["Docker Desktop", "obligatoria", "base de datos, API y servidor web"],
        ["LM Studio", "para el chat", "corre en el host: necesita la GPU"],
        ["uv", "para tests y scripts", "gestor de entorno Python"],
    ])
    p(d, "Sin el runtime de modelos la aplicación levanta igual y busca por "
         "texto; lo que no hay es respuesta redactada ni búsqueda semántica.")

    h1(d, "Instalación paso a paso")
    numerada(d, [
        "Clonar el repositorio y entrar al directorio.",
        "Instalar LM Studio y descargar dos modelos: el de chat y el de "
        "embeddings de 768 dimensiones.",
        "Dejar ambos residentes en memoria con un TTL largo. Sin esto cada "
        "consulta cuesta 3,8 s extra por recarga.",
        "Levantar el stack de contenedores.",
        "Cargar datos: o el volcado de demostración, que ya trae los "
        "embeddings calculados, o un scrapeo desde cero de cinco tiendas.",
    ])

    h1(d, "Comandos")
    tabla(d, ["Comando", "Qué hace"], [
        ["make up", "levanta base, API y web"],
        ["make down", "detiene el stack conservando los datos"],
        ["make modelo", "deja los dos modelos residentes"],
        ["make restore", "carga el volcado de demostración"],
        ["make dump", "exporta el conjunto actual comprimido"],
        ["make seed", "scrapea cinco tiendas desde cero"],
        ["make reset", "BORRA la base y la deja vacía"],
        ["make test", f"corre los {CIFRAS['tests']} tests de regresión"],
        ["make estado", "muestra qué hay cargado"],
    ])

    h1(d, "El volcado incluye los embeddings")
    p(d, "El archivo de volcado viaja con la columna de vectores ya calculada. "
         "Quien lo reciba no necesita el modelo ni volver a scrapear: levanta "
         "el stack, restaura, y la búsqueda semántica funciona.")

    h2(d, "Dos detalles que costaron una tarde")
    bullets(d, [
        "El volcado de datos escribe las tablas en orden alfabético, así que "
        "los productos entran antes que las tiendas y la clave foránea los "
        "rechaza. Se resuelve suspendiendo los triggers de integridad durante "
        "la carga.",
        "Un volcado con identificadores explícitos no mueve las secuencias de "
        "identidad: el siguiente INSERT arranca de 1 y choca. Hay que "
        "adelantarlas al máximo cargado.",
    ])

    h1(d, "Advertencia sobre el reinicio")
    p(d, "El comando de reinicio destruye el volumen de datos: se lleva puestos "
         "los productos, los embeddings y el histórico de precios. Ofrece hacer "
         "un volcado antes.")
    p(d, "El procedimiento seguro es: volcar, reiniciar, levantar, restaurar.")

    h1(d, "Trampa operativa: el contenedor hornea el código")
    p(d, "El servicio de API monta solo el directorio de datos; el código "
         "Python está dentro de la imagen. Reiniciar el contenedor corre la "
         "versión vieja. Después de tocar el código fuente hay que reconstruir "
         "la imagen. El directorio web sí está montado y se refleja recargando.")

    h1(d, "Cuando algo no anda")
    tabla(d, ["Síntoma", "Causa probable"], [
        ["El indicador del encabezado está en rojo",
         "el runtime de modelos no responde"],
        ["La primera pregunta tarda 4 s y las siguientes 1,4 s",
         "los modelos no quedaron residentes"],
        ["El chat contesta vacío",
         "presupuesto de tokens agotado por la cadena de razonamiento"],
        ["Busca pero no comprende",
         "el modelo de embeddings no responde; cayó a búsqueda por texto"],
        ["Dejó de encontrar tras cambiar el modelo de embeddings",
         "los vectores guardados son de otro modelo; hay que recalcularlos"],
    ])
    return d


# ─────────────────────────────────────────────────────────────────────────
# 10 · Estado y deudas
# ─────────────────────────────────────────────────────────────────────────
def doc10() -> Document:
    d = doc_nuevo(
        "FadIA · Estado actual, deudas y próximos pasos",
        "Qué está hecho, qué está roto y qué falta.",
    )

    h1(d, "Estado")
    tabla(d, ["Métrica", "Valor"], [
        ["Productos", CIFRAS["productos"]],
        ["Marcas", CIFRAS["marcas"]],
        ["Variantes", CIFRAS["variantes"]],
        ["Puntos de precio", CIFRAS["price_points"]],
        ["Cobertura de embeddings", "100 %"],
        ["Productos con imagen", "149.579"],
        ["Tests en verde", f"{CIFRAS['tests']} de {CIFRAS['tests']}"],
    ])

    h1(d, "Deudas técnicas conocidas")

    h2(d, "1. Normalización horneada al scrapear")
    p(d, "La clasificación corre solo en los adapters y queda congelada en el "
         "archivo intermedio. Arreglar el clasificador no corrige lo guardado. "
         "Dos consecuencias medidas: 3.514 productos de calzado clasificados "
         "como pantalón y 2.772 variantes con talles mal formateados.")
    p(d, "Solución: un pase de reclasificación offline. No requiere volver a "
         "scrapear.")

    h2(d, "2. El texto del embedding incluye lo que no debería")
    p(d, "Se embeben talles y colores, que además son filtros duros en SQL. "
         "Ahí solo diluyen: una consulta que dice «negro» deriva hacia "
         "productos que apenas listan negro entre seis colores. La marca "
         "aparece dos veces, lo que hace competir la similitud de nombre de "
         "marca con la de la prenda.")
    p(d, "Solución: limpiar el texto y recalcular. Conviene hacerlo junto con "
         "la reclasificación para no re-embeber dos veces.")

    h2(d, "3. Cobertura de marcas incompleta")
    p(d, "Alrededor de 30 marcas argentinas siguen sin adapter. La vía "
         "genérica por JSON-LD se midió y no alcanza: solo 2 de 25 publican el "
         "tipo Product. Harían falta selectores inducidos o un navegador "
         "headless.")

    h1(d, "Lo que no existe y no debería inventarse")
    bullets(d, [
        "Envíos y devoluciones. Ninguna de las siete plataformas los publica "
        "estructurados. Los 27.817 productos que mencionan «envío» lo tienen "
        "como texto libre en la descripción.",
        "Histórico de precios significativo. Solo 743 productos de "
        f"{CIFRAS['productos']} tienen un cambio real de precio. Un gráfico "
        "sería plano en el 99,5 % de los casos.",
    ])

    h1(d, "Oportunidades ordenadas por relación valor/esfuerzo")
    numerada(d, [
        "Reclasificación offline. Barata, arregla datos rotos hoy.",
        "Limpiar el texto del embedding. Una línea de cambio más una corrida "
        "de recálculo.",
        "Embeddings de imagen. El modelo de visión comparte espacio latente "
        "con el de texto, así que no hace falta recalcular el corpus. Pagaría "
        "sobre todo en el 29 % de productos sin descripción usable.",
        "Ampliar cobertura de marcas con navegador headless.",
        "Ingesta desde emails con planillas y fotos, y desde catálogos de "
        "Instagram y WhatsApp, para las marcas sin ecommerce.",
    ])

    h1(d, "Lo que hace a este proyecto difícil de copiar")
    bullets(d, [
        "El vocabulario textil argentino (morley, frisa, bengalina) extraído "
        "de catálogos reales.",
        "La serie temporal de precios, que se construye observando todos los "
        "días y no se puede comprar.",
        "El catálogo de fallas silenciosas codificado en el arnés de "
        "validación: cada test es un incidente real que ya pasó.",
        "La enumeración completa del marketplace mayorista, resuelta por "
        "sitemap después de comprobar que los listados se randomizan.",
    ])
    return d


DOCS = [
    ("01-que-es-fadia.docx", doc01),
    ("02-arquitectura.docx", doc02),
    ("03-adapters.docx", doc03),
    ("04-datos-y-normalizacion.docx", doc04),
    ("05-busqueda-y-memoria.docx", doc05),
    ("06-modelo-de-lenguaje.docx", doc06),
    ("07-fallas-silenciosas.docx", doc07),
    ("08-interfaz.docx", doc08),
    ("09-instalacion-y-operacion.docx", doc09),
    ("10-estado-y-deudas.docx", doc10),
]


def main() -> None:
    SALIDA.mkdir(parents=True, exist_ok=True)
    for nombre, fabricar in DOCS:
        destino = SALIDA / nombre
        fabricar().save(destino)
        print(f"  {destino}  ({destino.stat().st_size // 1024} KB)")
    print(f"\n{len(DOCS)} documentos en {SALIDA}/")


main()
